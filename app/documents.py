"""Documents added to a conversation (3.23): "Bu belgeyi özetle", "Belgede teslim tarihi geçiyor mu?".

PDF, Word (.docx) and plain text files are read here, on this computer; the text goes only to the local model.
Turkish and English documents both work: the model is told the document's language and answers in the language the
person asks in (an English contract can be asked about in Turkish, and the other way round).

How much the model can read at once is limited, so:
- a short document (up to FULL_CHARS, about 10-15 pages) goes whole into the conversation's system prompt: it stays
  the same for every question, so Ollama reads it once and later questions are fast;
- a longer one is searched for the passages that best match each question (by shared word stems, Turkish endings
  ignored), and "özetle" / "summarize" reads it part by part and summarizes the parts' summaries. The part summaries
  are kept, so a second summary question is quick.
"""

import json
import math
import re
from pathlib import Path

from . import db, llm

MAX_BYTES = 30 * 1024 * 1024
KINDS = {".pdf": "PDF", ".docx": "Word", ".txt": "metin", ".md": "metin", ".csv": "tablo (CSV)"}
DOC_CTX = 16384        # tokens the model reads in a conversation with documents (normal chats: llm.NUM_CTX)
FULL_CHARS = 28000     # all documents of a conversation up to this size go whole into the system prompt
PASSAGE_CHARS = 1500   # a long document is searched in pieces of about this size
PASSAGES_BUDGET = 18000
SECTION_CHARS = 18000  # a long document is summarized in parts of this size
MAX_SECTIONS = 24      # about 430 000 characters (~150 pages); the rest is not summarized
MAX_TEXT = 2_000_000


class DocumentError(Exception):
    pass


class ScannedPDF(DocumentError):
    """A PDF with pictures of pages only: read by the model page by page (pictures.read_scanned_pdf, 3.32)."""


# Reading the file

def extract(path: str, name: str) -> tuple[str, int | None]:
    """(text, number of pages or None)."""
    kind = Path(name).suffix.lower()
    if kind not in KINDS:
        raise DocumentError("Bu dosya türünü okuyamıyorum. PDF, Word (.docx), .txt, .md ya da .csv dosyası ekleyebilirsin."
                            + (" Eski Word (.doc) dosyasını Word'de açıp \"Farklı kaydet → .docx\" yapabilirsin."
                               if kind == ".doc" else ""))
    if kind == ".pdf":
        text, pages = _pdf(path)
    elif kind == ".docx":
        text, pages = _docx(path), None
    else:
        text, pages = _plain(path), None
    text = _tidy(text)
    if not text.strip():
        if kind == ".pdf":
            raise ScannedPDF("Bu PDF'te seçilebilir yazı yok: taranmış bir belge.")
        raise DocumentError("Dosyada okunacak yazı bulamadım.")
    return text[:MAX_TEXT], pages


def _pdf(path: str) -> tuple[str, int]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise DocumentError("Bu PDF şifreli. Şifresini kaldırıp (ya da yazdır → PDF olarak kaydet) tekrar ekle.")
        pages = []
        for number, page in enumerate(reader.pages, 1):
            try:
                pages.append(f"[Sayfa {number}]\n{page.extract_text() or ''}")
            except Exception:
                pages.append(f"[Sayfa {number}]")
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError(f"PDF okunamadı: {e}")
    body = "\n\n".join(pages)
    if not re.sub(r"\[Sayfa \d+\]", "", body).strip():
        return "", len(pages)
    return body, len(pages)


def _docx(path: str) -> str:
    import docx

    try:
        document = docx.Document(path)
    except Exception as e:
        raise DocumentError(f"Word dosyası okunamadı: {e}")
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:  # tables come after the text, row by row
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(dict.fromkeys(c for c in cells if c)))
    return "\n".join(parts)


def _plain(path: str) -> str:
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "cp1254", "latin-1"):  # cp1254: Turkish Windows
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # words split at the end of a line in PDFs
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# Language

_TR = {"ve", "bir", "bu", "için", "ile", "olarak", "da", "de", "olan", "gibi", "daha", "çok", "ne", "her", "ise"}
_EN = {"the", "and", "of", "to", "in", "is", "for", "that", "with", "on", "as", "are", "be", "this", "by"}


def language(text: str) -> str:
    sample = text[:20000].replace("I", "ı").replace("İ", "i").lower()
    words = re.findall(r"[a-zçğıöşü]+", sample)
    tr = sum(w in _TR for w in words) + 3 * len(re.findall(r"[çğış]", sample)) / 10
    en = sum(w in _EN for w in words)
    if tr + en < 5:
        return ""
    return "Türkçe" if tr >= en else "İngilizce"


def describe(doc: dict) -> str:
    """"12 sayfa, İngilizce" for the chat and the list."""
    parts = []
    if doc.get("pages"):
        parts.append(f"{doc['pages']} sayfa")
    chars = doc.get("chars") or len(doc.get("text") or "")
    parts.append(f"yaklaşık {max(1, round(chars / 6.5 / 100) * 100) if chars > 650 else max(1, round(chars / 6.5))} kelime")
    if doc.get("language"):
        parts.append(doc["language"])
    return ", ".join(parts)


# What the model gets

def fits_whole(docs: list[dict]) -> bool:
    return sum(len(d["text"]) for d in docs) <= FULL_CHARS


def prompt_block(docs: list[dict]) -> str:
    """Added to the conversation's system prompt; the same for every question, so Ollama's reading stays cached."""
    if not docs:
        return ""
    whole = fits_whole(docs)
    lines = ["", "Bu sohbete kullanıcı belge ekledi. Belgeyle ilgili sorularda YALNIZCA belgede yazanlara dayan; belgede",
             "olmayan bir şeyi uydurma, \"belgede bu bilgi yok\" de. Kullanıcı hangi dilde soruyorsa o dilde cevap ver",
             "(belge İngilizce olsa bile Türkçe sorulursa Türkçe). Alıntı yaparken belgenin kendi cümlesini ver ve gerekirse",
             "çevirisini ekle. Sayfa numarası biliniyorsa söyle ([Sayfa N] işaretleri)."]
    for doc in docs:
        title = f"=== BELGE: {doc['name']} ({describe(doc)}) ==="
        if whole:
            lines += ["", title, doc["text"], "=== BELGE SONU ==="]
        else:
            lines += ["", title, "(Uzun belge: tamamı burada değil. Her soruda belgenin soruyla ilgili bölümleri, kullanıcının",
                      "mesajının sonunda \"Belgeden ilgili bölümler\" olarak verilir; yalnızca onlara dayan.)",
                      "Belgenin başı:", doc["text"][:1500], "=== ==="]
    return "\n".join(lines)


_STOP = {"bir", "bu", "şu", "ve", "ile", "için", "ne", "neler", "nedir", "var", "mı", "mi", "mu", "mü", "da", "de",
         "belge", "belgede", "belgenin", "belgeyi", "dosya", "dosyada", "the", "and", "what", "does", "is", "are",
         "of", "to", "in", "document", "file", "hangi", "nasıl", "kaç", "nerede", "söyle", "anlat", "geçiyor"}
_FOLD = str.maketrans("çğıöşüâîû", "cgiosuaiu")


def _stems(text: str) -> list[str]:
    low = text.replace("I", "ı").replace("İ", "i").lower()
    return [w.translate(_FOLD)[:5] for w in re.findall(r"[\wçğıöşü]+", low) if len(w) >= 3 and w not in _STOP]


def _pieces(text: str) -> list[str]:
    """Pieces of about PASSAGE_CHARS, cut at paragraph or sentence ends."""
    pieces, current = [], ""
    for para in re.split(r"\n\s*\n", text):
        while len(para) > PASSAGE_CHARS:
            cut = max(para.rfind(". ", 0, PASSAGE_CHARS), para.rfind("\n", 0, PASSAGE_CHARS))
            cut = cut + 1 if cut > PASSAGE_CHARS // 3 else PASSAGE_CHARS
            if current:
                pieces.append(current)
                current = ""
            pieces.append(para[:cut].strip())
            para = para[cut:]
        if len(current) + len(para) > PASSAGE_CHARS and current:
            pieces.append(current)
            current = ""
        current = (current + "\n\n" + para).strip()
    if current:
        pieces.append(current)
    return pieces


def passages(docs: list[dict], question: str) -> str:
    """The parts of long documents that best match the question, in document order."""
    wanted = set(_stems(question))
    found = []  # (score, doc index, piece index, doc name, piece)
    all_pieces = [(i, j, doc["name"], piece) for i, doc in enumerate(docs) for j, piece in enumerate(_pieces(doc["text"]))]
    if not all_pieces:
        return ""
    counts = [_stems(p[3]) for p in all_pieces]
    for (i, j, name, piece), stems in zip(all_pieces, counts):
        score = 0.0
        for stem in wanted:
            hits = stems.count(stem)
            if hits:
                have = sum(stem in c for c in counts)
                score += (1 + math.log(hits)) * math.log(1 + len(counts) / have)
        found.append((score, i, j, name, piece))
    best = sorted((f for f in found if f[0] > 0), key=lambda f: -f[0])
    if not best:  # nothing matches (e.g. "ne anlatıyor?"): the start of each document
        best = sorted(found, key=lambda f: (f[2], f[1]))
    chosen, used = [], 0
    for item in best:
        if used + len(item[4]) > PASSAGES_BUDGET:
            continue
        chosen.append(item)
        used += len(item[4])
    chosen.sort(key=lambda f: (f[1], f[2]))
    return "\n\n".join(f"--- {name}, bölüm {j + 1} ---\n{piece}" for _, _, j, name, piece in chosen)


_SUMMARY = re.compile(r"\b(özetle|özetler|özet|özetini|özetin|ne anlatıyor|neyi anlatıyor|ne hakkında|konusu ne|"
                      r"ana fikr|genel olarak|summari[sz]e|summary|overview|what is it about|tl;?dr)", re.IGNORECASE)


def wants_summary(question: str) -> bool:
    return bool(_SUMMARY.search(question))


def _sections(text: str) -> list[str]:
    sections, current = [], ""
    for piece in _pieces(text):
        if len(current) + len(piece) > SECTION_CHARS and current:
            sections.append(current)
            current = ""
        current = (current + "\n\n" + piece).strip()
    if current:
        sections.append(current)
    return sections


def question_language(question: str) -> str:
    low = question.replace("I", "ı").replace("İ", "i").lower()
    words = set(re.findall(r"[a-zçğıöşü]+", low))
    return "Türkçe" if re.search(r"[çğıöşü]", low) or words & (_TR | {"nedir", "var", "mı", "mi", "ne", "kaç"}) else "İngilizce"


async def search_words(model: str, docs: list[dict], question: str) -> str:
    """The question's key words in the documents' language too, so a Turkish question finds an English passage."""
    asked = question_language(question)
    other = {d.get("language") for d in docs} - {asked, "", None}
    if not other:
        return question
    lang = "English" if "İngilizce" in other else "Turkish"
    try:
        words = await llm.chat_text(model, [{"role": "user", "content":
            f"Translate the important search words of this question into {lang}, including likely synonyms. "
            f"Reply with the words only, separated by spaces, nothing else.\n\nQuestion: {question}"}],
            num_ctx=DOC_CTX, num_predict=40)
    except llm.OllamaError:
        return question
    return f"{question} {words}"


async def section_summaries(model: str, doc: dict):
    """Summaries of a long document's parts (made once, then kept).

    Yields ("progress", text) while working and finally ("done", (summaries, whole document covered?)).
    """
    sections = _sections(doc["text"])
    covered = len(sections) <= MAX_SECTIONS
    sections = sections[:MAX_SECTIONS]
    if doc.get("summary"):
        try:
            kept = json.loads(doc["summary"])
            if len(kept) == len(sections):
                yield "done", (kept, covered)
                return
        except ValueError:
            pass
    summaries = []
    for number, section in enumerate(sections, 1):
        yield "progress", f"📄 {doc['name']}: bölüm {number}/{len(sections)} okunuyor…"
        summaries.append(await llm.chat_text(model, [
            {"role": "system", "content": "Bir belgenin bir bölümünü özetliyorsun. Önemli bilgileri (kişiler, tarihler, "
                                          "tutarlar, kararlar, yükümlülükler) kaçırmadan kısa maddelerle özetle. "
                                          "Belge hangi dildeyse özeti o dilde yaz. Yorum ekleme."},
            {"role": "user", "content": f"Belge: {doc['name']}, bölüm {number}/{len(sections)}:\n\n{section}"},
        ], num_ctx=DOC_CTX))
    db.set_document_summary(doc["id"], json.dumps(summaries, ensure_ascii=False))
    yield "done", (summaries, covered)
